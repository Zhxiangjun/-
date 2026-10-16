# -*- coding: utf-8 -*-
"""交付物一致性审计：核对主数据中的关键数字在各文档中的出现情况。"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from feature_master import build, MODULES, stage_summary
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

BASE = r"F:\职业中专\大赛\智慧园区创意"
rows = build()
st = stage_summary(rows)
TOTAL = sum(r["effort"] for r in rows)
N = len(rows)

EXPECT = {
    "功能项总数": str(N),
    "模块数": str(len(MODULES)),
    "总人天": str(TOTAL),
    "MVP项数": str(st["MVP"]["count"]),
    "MVP人天": str(st["MVP"]["effort"]),
    "V1项数": str(st["V1"]["count"]),
    "V1人天": str(st["V1"]["effort"]),
    "V2项数": str(st["V2"]["count"]),
    "V2人天": str(st["V2"]["effort"]),
    "V3项数": str(st["V3"]["count"]),
    "V3人天": str(st["V3"]["effort"]),
}


def docx_text(p):
    d = Document(p)
    t = "\n".join(x.text for x in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                t += "\n" + c.text
    return t


def pptx_text(p):
    prs = Presentation(p)
    t = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                t.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        t.append(c.text)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for pl in sh.chart.plots:
                    for srs in pl.series:
                        t.append(str(srs.name))
                        t.append(" ".join(str(v) for v in srs.values))
                t.extend([str(c) for c in sh.chart.plots[0].categories])
    return "\n".join(t)


def xlsx_text(p):
    wb = load_workbook(p, data_only=False)
    t = []
    for ws in wb.worksheets:
        t.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            t.append(" ".join(str(v) for v in row if v is not None))
    return "\n".join(t)


FILES = [
    ("01_需求规格说明书.docx", docx_text),
    ("02_总体技术架构方案.docx", docx_text),
    ("03_功能清单与优先级评估表.xlsx", xlsx_text),
    ("04_实施路线图与里程碑计划.docx", docx_text),
    ("05_商业计划书.docx", docx_text),
    ("06_项目汇报PPT.pptx", pptx_text),
]

print(f"主数据：{len(MODULES)} 模块 / {N} 项功能 / {TOTAL} 人天")
print(f"阶段分布：MVP {st['MVP']['count']}项{st['MVP']['effort']}人天 ｜ V1 {st['V1']['count']}项{st['V1']['effort']}人天 ｜ "
      f"V2 {st['V2']['count']}项{st['V2']['effort']}人天 ｜ V3 {st['V3']['count']}项{st['V3']['effort']}人天")
print()
texts = {}
for fn, reader in FILES:
    p = os.path.join(BASE, fn)
    if not os.path.exists(p):
        print("缺失:", fn)
        continue
    texts[fn] = reader(p)
    print(f"{fn:<32} 提取字符数 {len(texts[fn]):>7}")

print("\n关键数字出现情况（√=出现，- =未出现）：")
hdr = f"{'指标':<12}" + "".join(f"{fn.split('_')[0]:>7}" for fn, _ in FILES)
print(hdr)
for k, v in EXPECT.items():
    line = f"{k:<12}"
    for fn, _ in FILES:
        t = texts.get(fn, "")
        line += f"{'√' if v in t else '-':>7}"
    print(line)

print("\n跨文档金额口径核对：")
KEY_MONEY = {
    "人力成本合计（元）": f"{TOTAL * 1200:,}",
    "MVP人力成本（元）": f"{st['MVP']['effort'] * 1200:,}",
    "MVP+V1人力成本（元）": f"{(st['MVP']['effort'] + st['V1']['effort']) * 1200:,}",
}
for k, v in KEY_MONEY.items():
    hits = [fn for fn, _ in FILES if v in texts.get(fn, "")]
    print(f"  {k}: {v} → 出现在 {hits if hits else '（无，属正常，不同文档按万元/阶段口径表述）'}")
