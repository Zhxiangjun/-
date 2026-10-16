# -*- coding: utf-8 -*-
"""量化环境类文档之间的文本重合度（5-gram Jaccard），给出客观重复度。"""
import os
import re
import itertools
from docx import Document
from openpyxl import load_workbook

BASE = r"F:\职业中专\大赛\智慧园区创意\项目文档"


def docx_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def xlsx_text(path):
    wb = load_workbook(path, data_only=False)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.append(" | ".join(str(v) for v in row if v is not None))
    return "\n".join(parts)


FILES = [
    ("07准备", "07_开发环境搭建与部署准备指南.docx", docx_text),
    ("08核对", "08_环境准备与选型核对表.xlsx", xlsx_text),
    ("09方案V2", "09_电脑环境配置详细方案.docx", docx_text),
    ("10台账V2", "10_环境配置进度与验收台账（单机版）.xlsx", xlsx_text),
    ("10台账V1", "10_环境配置进度与验收台账.xlsx", xlsx_text),
    ("11总览", "11_电脑环境配置总览与快速开始.docx", docx_text),
    ("12协作", "12_协作开发环境配置指引.docx", docx_text),
]


def grams(text, n=5):
    t = re.sub(r"\s+", "", text)
    t = re.sub(r"[|·•\-—＝=*#]+", "", t)
    return {t[i:i + n] for i in range(max(0, len(t) - n + 1))}


G = {}
for label, name, reader in FILES:
    p = os.path.join(BASE, name)
    if os.path.exists(p):
        G[label] = grams(reader(p))

print("=" * 96)
print("环境类文档两两重合度（5-gram Jaccard；≥40% 属显著重复，≥65% 属近乎重复）")
print("=" * 96)
labels = list(G.keys())
print(f"{'':<10}" + "".join(f"{l:>10}" for l in labels))
for a in labels:
    row = f"{a:<10}"
    for b in labels:
        if a == b:
            row += f"{'—':>10}"
        else:
            inter = len(G[a] & G[b])
            union = len(G[a] | G[b])
            row += f"{inter / union * 100:>9.1f}%"
    print(row)

print()
print("=" * 96)
print("重点结论")
print("=" * 96)
pairs = []
for a, b in itertools.combinations(labels, 2):
    inter = len(G[a] & G[b])
    union = len(G[a] | G[b])
    pairs.append((inter / union * 100, a, b))
pairs.sort(reverse=True)
for pct, a, b in pairs:
    tag = "近乎重复" if pct >= 65 else ("显著重复" if pct >= 40 else ("轻度重叠" if pct >= 20 else "基本独立"))
    print(f"  {a:<10} ↔ {b:<10} {pct:>5.1f}%   {tag}")
