# -*- coding: utf-8 -*-
"""生成 08_环境准备与选型核对表.xlsx"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"F:\职业中专\大赛\智慧园区创意\08_环境准备与选型核对表.xlsx"

NAVY, BLUE, LIGHT, GREY = "1F3864", "2E5C8A", "DCE6F1", "F2F2F2"
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row=1, fill=NAVY):
    for c in ws[row]:
        if c.value is None:
            continue
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def style_body(ws, first=2, last=None, size=9):
    last = last or ws.max_row
    for r in ws.iter_rows(min_row=first, max_row=last, max_col=ws.max_column):
        for c in r:
            c.font = Font(name="微软雅黑", size=size)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name="微软雅黑", size=14, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1, value=sub).font = Font(name="微软雅黑", size=9, color="808080")
    return 3


wb = Workbook()

# ============================================================
# 01 机器角色分工
# ============================================================
ws = wb.active
ws.title = "01-机器角色分工"
title(ws, "开发环境机器角色分工表", "版本 V1.0 ｜ 2026 年 11 月 ｜ 依据团队现有三台电脑 + 阿里云 ECS 的实际配置", 8)
hdr = ["机器", "CPU", "内存", "显卡", "系统选择", "承担角色", "必装组件", "不要安装/不要做"]
ws.append([])
r = 4
for i, h in enumerate(hdr, start=1):
    ws.cell(row=r, column=i, value=h)
style_header(ws, r)
DATA = [
    ["台式机（主力）", "i5-13400F（10 核 16 线程）", "16GB（建议升 32GB）", "RTX 3050 8GB（算力 8.6）",
     "Windows 11 + WSL2 Ubuntu 24.04",
     "AI 训练与 ONNX 导出 + 全栈一体化开发 + Docker 起依赖服务",
     "Git、VS Code、Python 3.12 与 3.11、Node LTS+pnpm、Docker Desktop、DBeaver、MQTTX、Apifox、OBS、CUDA+PyTorch+Ultralytics",
     "不要作为唯一数据保存地；不要把数据库、样本、权重全放系统盘"],
    ["笔记本 1", "AMD Ryzen 7 5800U（8 核 16 线程）", "16GB", "Radeon 集显（496MB，无 CUDA）",
     "Windows 11 + WSL2 Ubuntu 24.04",
     "服务宿主与演示机：Docker 起 PG/Redis/EMQX/MinIO + 后端 + 前端",
     "Git、VS Code、Python 3.12、Node+pnpm、Docker Desktop、MQTTX、Apifox、MobaXterm、OBS",
     "不要做模型训练；Docker 内存分配不要超过 10GB"],
    ["笔记本 2", "i5-7300HQ（4 核 4 线程）", "16GB", "GTX 1050 2GB（算力 6.1）",
     "Windows 10/11（保持原系统）",
     "客户端与设备联调：多角色浏览器测试、串口/网口调试、现场备机",
     "Git、VS Code、Node+pnpm、MQTTX、SSCOM/XCOM、ONVIF Device Manager、MobaXterm",
     "不要装 Docker 全栈；不要折腾 CUDA 训练环境；不要当数据库服务器"],
    ["阿里云 ECS", "2 核（起步）→ 4 核（规模）", "4GB（起步）→ 8GB（规模）", "无",
     "Ubuntu 24.04 LTS",
     "业务中枢与对外发布：API、数据库、报表、小程序后端、异地备份",
     "Docker + Compose v2、Nginx、Certbot、Git、备份脚本、云监控",
     "不要接收校内摄像头视频流；数据库端口不要对公网开放"],
]
for i, row in enumerate(DATA, start=1):
    for j, v in enumerate(row, start=1):
        ws.cell(row=r + i, column=j, value=v)
style_body(ws, r + 1, r + len(DATA))
widths(ws, [16, 22, 16, 24, 26, 32, 40, 34])
ws.freeze_panes = "A5"

# ============================================================
# 02 软件安装清单
# ============================================================
ws = wb.create_sheet("02-软件安装清单")
title(ws, "软件安装与配置清单（可勾选核对）", "用途栏标注适用机器；安装顺序建议自上而下；安装完在“状态”列标记 已装/未装/不适用", 9)
hdr = ["序号", "类别", "软件/组件", "建议版本", "适用机器", "用途", "获取方式", "关键配置", "状态"]
ws.append([])
r = 4
for i, h in enumerate(hdr, start=1):
    ws.cell(row=r, column=i, value=h)
style_header(ws, r)
SW = [
    (1, "基础", "Git for Windows", "最新稳定版", "全部", "版本管理与多机同步", "git-scm.com", "设置 user.name/user.email；开启 core.autocrlf=false", ""),
    (2, "基础", "VS Code", "最新稳定版", "全部", "主开发工具", "code.visualstudio.com", "安装 Python/Vue/Ruff/Docker/Remote-WSL 扩展", ""),
    (3, "基础", "Windows Terminal + PowerShell 7", "最新版", "Windows 机器", "命令行体验与 UTF-8 兼容", "Microsoft Store / winget", "设置默认终端为 Windows Terminal", ""),
    (4, "基础", "Python 3.12", "3.12.x", "台式机/笔记本1", "后端开发", "python.org", "勾选 Add to PATH；设置 PYTHONUTF8=1", ""),
    (5, "基础", "Python 3.11", "3.11.x", "台式机", "AI 训练生态兼容", "python.org", "与 3.12 并存，用虚拟环境隔离", ""),
    (6, "基础", "Node.js LTS", "22.x LTS", "全部", "前端构建", "nodejs.org", "启用 corepack；npm 源切 npmmirror", ""),
    (7, "基础", "pnpm", "最新稳定版", "全部", "前端包管理", "corepack enable", "corepack prepare pnpm@latest --activate", ""),
    (8, "基础", "uv", "最新版", "台式机/笔记本1", "Python 依赖与虚拟环境（比 pip 快很多）", "astral.sh/uv", "配置清华 pip 源", ""),
    (9, "容器", "WSL2 + Ubuntu 24.04", "Ubuntu 24.04 LTS", "台式机/笔记本1", "Docker 后端与 Linux 部署环境", "wsl --install -d Ubuntu-24.04", "BIOS 开启 VT-x/SVM；apt 源换阿里云", ""),
    (10, "容器", "Docker Desktop", "最新稳定版", "台式机/笔记本1", "一键起 PG/Redis/EMQX/MinIO", "docker.com", "勾选 WSL2 后端；配置阿里云 ACR 镜像加速；限制内存", ""),
    (11, "数据库", "DBeaver Community", "最新版", "全部", "连接与调试 PostgreSQL", "dbeaver.io", "创建 campus 连接；保存常用 SQL", ""),
    (12, "数据库", "PostgreSQL 16 镜像", "postgres:16 / pgvector:pg16", "台式机/笔记本1/云", "业务主库", "docker pull", "强口令；数据卷持久化；TZ=Asia/Shanghai", ""),
    (13, "数据库", "Redis 7 镜像", "redis:7-alpine", "台式机/笔记本1/云", "缓存与白名单", "docker pull", "启用 AOF 与密码", ""),
    (14, "物联网", "EMQX 5 镜像", "emqx/emqx:5", "台式机/笔记本1/云", "MQTT 设备接入", "docker pull", "首次登录改默认密码；规划主题命名", ""),
    (15, "对象存储", "MinIO 镜像", "minio/minio", "台式机/笔记本1/云", "事件截图与报表存储", "docker pull", "私有桶；强密钥；生命周期策略", ""),
    (16, "AI 训练", "NVIDIA 显卡驱动", "官方最新 Studio/Game Ready", "台式机", "CUDA 支持", "nvidia.com", "先装驱动，再装 PyTorch", ""),
    (17, "AI 训练", "PyTorch（cu12x）", "与驱动匹配版本", "台式机", "模型训练", "pytorch.org 生成的安装命令", "用独立 venv；验证 torch.cuda.is_available()", ""),
    (18, "AI 训练", "Ultralytics（YOLO）", "最新版", "台式机", "检测模型训练与导出", "pip install ultralytics", "准备 data.yaml；先用官方小模型跑通", ""),
    (19, "AI 训练", "ONNX Runtime + onnxsim", "最新版", "台式机", "边缘推理与模型精简", "pip install onnxruntime onnxsim", "统一 opset 17；导出后做逐张比对", ""),
    (20, "AI 训练", "LabelImg / X-AnyLabeling", "最新版", "台式机", "数据标注", "pip install labelImg / GitHub", "先定标注规范：类别、最小框、遮挡", ""),
    (21, "AI 训练", "OpenCV、NumPy、Pillow", "最新稳定版", "台式机", "图像处理基础库", "pip install", "opencv-python 而非 opencv-contrib 除非需要", ""),
    (22, "前端", "Vite + Vue3 + TS 脚手架", "最新版", "全部", "管理端与驾驶舱", "pnpm create vite", "配置 /api 代理到 8000", ""),
    (23, "前端", "Element Plus / Pinia / Vue Router / Axios", "最新版", "全部", "UI、状态、路由、请求", "pnpm add", "Axios 统一拦截器；按角色划分路由", ""),
    (24, "前端", "ECharts", "5.x", "全部", "统计图表与驾驶舱大屏", "pnpm add echarts", "1920×1080 栅格自适应缩放", ""),
    (25, "前端", "uni-app", "最新版", "按需", "学生/教师/家长小程序", "HBuilderX 或 CLI", "配置小程序 appid 与服务器域名白名单", ""),
    (26, "设备模拟", "MQTTX", "最新版", "全部", "模拟设备上报与订阅", "mqttx.app", "建立本地/云端连接配置", ""),
    (27, "设备模拟", "FFmpeg", "最新稳定版", "台式机/笔记本1", "伪视频流与录像处理", "ffmpeg.org / apt", "配合 MediaMTX 推 RTSP 流", ""),
    (28, "设备模拟", "MediaMTX", "最新版", "台式机/笔记本1", "轻量 RTSP/RTMP 服务器", "docker pull bluenviron/mediamtx", "rtsp://localhost:8554/cam1", ""),
    (29, "设备模拟", "ONVIF Device Manager", "最新版", "笔记本2", "摄像机发现与参数调试", "GitHub", "用于现场快速定位摄像机 IP", ""),
    (30, "调试", "Apifox 或 Postman", "最新版", "全部", "接口联调与自动化测试", "官网", "导入 Swagger；维护测试集合", ""),
    (31, "调试", "MobaXterm / WinSCP", "最新版", "全部", "SSH 运维与文件传输", "官网", "保存云端与现场服务器会话", ""),
    (32, "调试", "SSCOM / XCOM", "最新版", "笔记本2", "串口调试（485/Modbus 设备）", "网络下载", "记录常用波特率与指令集", ""),
    (33, "调试", "Wireshark", "最新版", "台式机/笔记本2", "抓包排查设备协议", "wireshark.org", "抓门禁与视频协议报文时使用", ""),
    (34, "演示", "OBS Studio", "最新版", "台式机/笔记本1", "录制演示视频（比赛备份）", "obsproject.com", "按 1920×1080 录制一次完整流程", ""),
    (35, "云侧", "Nginx + Certbot", "最新稳定版", "云服务器", "反向代理与 HTTPS", "apt install", "证书自动续期；仅放行必要端口", ""),
    (36, "云侧", "Docker Engine + Compose v2", "最新稳定版", "云服务器/现场服务器", "容器编排", "get.docker.com", "配置镜像加速；限制容器资源", ""),
    (37, "云侧", "云监控与快照策略", "阿里云控制台", "云服务器", "可用性与数据保护", "控制台配置", "CPU/内存/磁盘/带宽告警到手机", ""),
    (38, "办公", "Office/WPS、Markdown 编辑器", "最新版", "全部", "文档编写与交付", "官网", "统一文档模板与命名规范", ""),
]
for i, row in enumerate(SW, start=1):
    for j, v in enumerate(row, start=1):
        ws.cell(row=r + i, column=j, value=v)
style_body(ws, r + 1, r + len(SW))
widths(ws, [6, 10, 30, 22, 16, 30, 26, 34, 10])
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{r}:I{r + len(SW)}"

# ============================================================
# 03 环境验收清单
# ============================================================
ws = wb.create_sheet("03-环境验收清单")
title(ws, "环境验收清单（逐项打勾）", "覆盖开发机、云服务器、机房、现场四类环境；验收不通过不得进入下一步开发或部署", 6)
hdr = ["序号", "环境", "验收项", "验证方法/命令", "通过标准", "状态"]
ws.append([])
r = 4
for i, h in enumerate(hdr, start=1):
    ws.cell(row=r, column=i, value=h)
style_header(ws, r)
CHECK = [
    (1, "台式机", "Python 双版本", "python --version / py -3.11 --version", "分别输出 3.12.x 与 3.11.x", ""),
    (2, "台式机", "Node 与 pnpm", "node -v && pnpm -v", "Node 为 LTS，pnpm 正常输出版本", ""),
    (3, "台式机", "WSL2", "wsl -l -v", "Ubuntu-24.04 状态 Running，VERSION 为 2", ""),
    (4, "台式机", "Docker", "docker run --rm hello-world", "输出 Hello from Docker!", ""),
    (5, "台式机", "依赖服务全起", "docker compose -f docker-compose.dev.yml ps", "5 个服务 healthy/running", ""),
    (6, "台式机", "数据库连通", "DBeaver 连 localhost:5432", "可看到 campus 库与表", ""),
    (7, "台式机", "MQTT 通", "MQTTX 连 1883 收发消息", "平台能收到设备消息", ""),
    (8, "台式机", "后端可跑", "访问 http://localhost:8000/docs", "Swagger 正常", ""),
    (9, "台式机", "前端可跑", "pnpm dev 后访问 5173", "页面正常且 /api 代理生效", ""),
    (10, "台式机", "GPU 可训", "torch.cuda.is_available()；跑 10 轮小训练", "输出 True 且训练完成生成权重", ""),
    (11, "台式机", "ONNX 导出与推理一致", "yolo export 后 onnxruntime 推理一张图", "与 PyTorch 偏差 ≤1e-3", ""),
    (12, "台式机", "伪视频可拉流", "平台拉 rtsp://localhost:8554/cam1", "画面正常，AI 能出结果", ""),
    (13, "台式机", "代码已入库", "git push 到 Gitee", "另一台机器能 clone 并运行", ""),
    (14, "笔记本1", "Docker 内存受限配置", "Docker Desktop → Resources", "分配 8—10GB，Windows 不卡顿", ""),
    (15, "笔记本1", "演示环境可离线运行", "断网后访问本机平台", "全部功能可用（除短信/微信）", ""),
    (16, "笔记本1", "演示录像备份", "OBS 录制一次完整流程", "视频 ≥10 分钟且画面清晰", ""),
    (17, "笔记本2", "多角色并发测试", "同时登录 5 类角色", "权限隔离正确，无越权可见数据", ""),
    (18, "笔记本2", "设备调试工具就绪", "MQTTX/串口/ONVIF 工具可用", "能连上至少一类真实或模拟设备", ""),
    (19, "云服务器", "SSH 安全加固", "仅密钥登录，禁用 root 直登", "口令登录被拒绝", ""),
    (20, "云服务器", "安全组最小开放", "控制台查看规则", "仅 22/80/443 对公网开放", ""),
    (21, "云服务器", "容器全栈运行", "docker compose ps", "全部 healthy", ""),
    (22, "云服务器", "HTTPS 有效", "浏览器访问域名", "无证书警告且自动续期已配置", ""),
    (23, "云服务器", "接口可访问", "访问 https://域名/docs", "正常，鉴权错误码符合预期", ""),
    (24, "云服务器", "备份可恢复", "从 OSS 备份恢复到测试库", "数据一致，流程可复现", ""),
    (25, "云服务器", "监控告警", "触发一次阈值告警", "手机收到告警通知", ""),
    (26, "云服务器", "时间与日志", "timedatectl / 日志目录", "北京时间为准，日志按天切割保留 ≥30 天", ""),
    (27, "机房", "母盘环境完成", "母盘机跑通全部教学用例", "学生按文档可独立完成一次提交", ""),
    (28, "机房", "批量部署完成", "随机抽 5 台机器检查", "环境一致，无缺件", ""),
    (29, "机房", "离线包可用", "断网机器上安装依赖", "无需外网即可装完", ""),
    (30, "机房", "开发服务器可用", "学生机访问统一服务器地址", "数据库/API 可连，账号隔离", ""),
    (31, "机房", "还原与修复策略", "执行一次环境修复脚本", "5 分钟内恢复可用环境", ""),
    (32, "现场", "服务器系统与磁盘", "cat /etc/os-release；df -h", "Ubuntu 24.04；数据盘已挂载且重启自动挂载", ""),
    (33, "现场", "时间同步", "chronyc sources", "偏差 ≤50ms", ""),
    (34, "现场", "网络与 VLAN", "三网隔离验证", "业务/视频/物联网互不干扰", ""),
    (35, "现场", "IP 台账", "核对设备台账表", "编号—位置—IP—MAC 完整无缺", ""),
    (36, "现场", "协议与端口确认", "与厂商逐项确认", "标准协议可用，端口已放行", ""),
    (37, "现场", "门禁联调", "现场刷卡/刷脸测试；断网 30 分钟", "P95 ≤300ms；离线放行且记录补传完整", ""),
    (38, "现场", "视频接入", "按点位清单逐个确认", "画面与录像正常，预览延时 ≤1.5 秒", ""),
    (39, "现场", "AI 算法验收", "用验收样本集测试", "准确率与误报率达标", ""),
    (40, "现场", "备份恢复演练", "从备份恢复一次", "RPO ≤15 分钟，RTO ≤2 小时", ""),
    (41, "现场", "培训与文档移交", "三类角色培训并签到", "签到记录与手册齐全", ""),
]
for i, row in enumerate(CHECK, start=1):
    for j, v in enumerate(row, start=1):
        ws.cell(row=r + i, column=j, value=v)
style_body(ws, r + 1, r + len(CHECK))
widths(ws, [6, 10, 26, 40, 44, 10])
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{r}:F{r + len(CHECK)}"

# ============================================================
# 04 阿里云选型与成本
# ============================================================
ws = wb.create_sheet("04-云服务器选型与成本")
title(ws, "阿里云 ECS 选型与年度成本测算", "价格为公开资料的量级估算（元/年，含税），实际以阿里云当期报价与优惠政策为准", 6)
hdr = ["项目", "起步配置（推荐先用）", "规模配置（小程序上线后）", "备注", "年费用估算(起步)", "年费用估算(规模)"]
ws.append([])
r = 4
for i, h in enumerate(hdr, start=1):
    ws.cell(row=r, column=i, value=h)
style_header(ws, r)
CLOUD = [
    ["实例规格", "2 核 4GB（经济型 e / 通用型 u1）", "4 核 8GB（通用型）", "新用户与学生认证有较大折扣", "300—800", "1500—3000"],
    ["操作系统", "Ubuntu 24.04 LTS", "同左", "与 WSL、现场服务器保持一致", "0", "0"],
    ["系统盘", "ESSD 40—60GB", "ESSD 100GB", "容器镜像与数据库占用较大", "100—200", "300—600"],
    ["公网带宽", "按使用流量（峰值 5—10Mbps）", "按量 + 峰值限速告警", "固定小带宽在拉镜像/传文件时会成为瓶颈", "100—400", "400—1200"],
    ["域名", ".cn 或学校提供的 edu.cn 域名", "同左", "小程序要求备案域名", "30—80", "30—80"],
    ["SSL 证书", "免费 DV 证书或 Let's Encrypt", "同左", "每年申请或自动续期", "0", "0"],
    ["OSS 对象存储", "40GB 标准存储 + 生命周期转低频", "200GB", "事件截图、报表、备份归档", "50—150", "300—800"],
    ["短信服务", "按条计费（未归预警、审批通知）", "按条计费", "需模板报备，注意发送频率限制", "0—200", "300—1000"],
    ["自动快照", "每周 1 次", "每日 1 次", "保护数据，避免误删不可恢复", "0—100", "100—300"],
    ["合计（量级估算）", "—", "—", "不含小程序开发与人力成本", "约 600—1900", "约 3000—7000"],
]
for i, row in enumerate(CLOUD, start=1):
    for j, v in enumerate(row, start=1):
        ws.cell(row=r + i, column=j, value=v)
style_body(ws, r + 1, r + len(CLOUD))
widths(ws, [18, 34, 28, 40, 18, 18])
ws.cell(row=r + len(CLOUD) + 1, column=1,
        value="提示：先用起步配置跑通全流程（多数新用户首年可控制在千元以内），待小程序对师生开放前再升级实例规格与带宽。").font = Font(name="微软雅黑", size=9, color="808080")

# ============================================================
# 05 机房离线包清单
# ============================================================
ws = wb.create_sheet("05-机房离线包清单")
title(ws, "学生机房标准化离线包清单", "用于断网或限网机房；建议用一块 32GB 以上移动固态统一存放并建立版本目录", 5)
hdr = ["序号", "类别", "内容", "获取方式", "体积参考"]
ws.append([])
r = 4
for i, h in enumerate(hdr, start=1):
    ws.cell(row=r, column=i, value=h)
style_header(ws, r)
OFF = [
    (1, "系统与驱动", "Ubuntu Server 24.04 ISO（现场服务器用）", "官网/校内镜像", "约 2.5GB"),
    (2, "基础软件", "Git、VS Code、Python 3.12、Node LTS、DBeaver、MobaXterm 安装包", "官网下载 Windows 版", "约 1.5—2GB"),
    (3, "Python 依赖", "后端与算法依赖 wheels（pip download -d wheels -r requirements.txt）", "在有网机器上预下载", "约 1—3GB"),
    (4, "前端依赖", "pnpm store（pnpm fetch 后整目录拷贝）", "pnpm fetch", "约 300—800MB"),
    (5, "编辑器扩展", "VS Code 扩展 .vsix（Python/Vue/Ruff/Docker/中文包）", "市场页下载", "约 200—500MB"),
    (6, "容器镜像", "docker save 导出的 postgres/redis/emqx/minio 镜像 tar", "docker save", "约 2—4GB"),
    (7, "数据与样本", "标注数据集、测试视频、模型权重", "ModelScope / Kaggle / 自行采集", "视规模而定"),
    (8, "安装脚本", "install-baseline.ps1、check-env.py、环境修复脚本", "项目组自研", "<1MB"),
    (9, "文档与课件", "本套项目文档、操作手册、作业模板", "项目仓库导出", "几十 MB"),
    (10, "演示素材", "演示视频、演示脚本、离线版前端构建产物", "OBS 录制/构建导出", "几百 MB"),
]
for i, row in enumerate(OFF, start=1):
    for j, v in enumerate(row, start=1):
        ws.cell(row=r + i, column=j, value=v)
style_body(ws, r + 1, r + len(OFF))
widths(ws, [6, 14, 52, 30, 16])

wb.save(OUT)
print("已生成:", OUT)
print("工作表:", wb.sheetnames)
