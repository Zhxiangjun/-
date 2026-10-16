# -*- coding: utf-8 -*-
import re
from docx import Document

for p in [r"F:\职业中专\大赛\智慧园区创意\02_总体技术架构方案.docx",
          r"F:\职业中专\大赛\智慧园区创意\05_商业计划书.docx"]:
    d = Document(p)
    hits = []
    for x in d.paragraphs:
        if "**" in x.text or "`" in x.text:
            hits.append(("P", x.text[:150]))
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                if "**" in c.text or "`" in c.text:
                    hits.append(("T", c.text[:150]))
    print("=====", p.split("\\")[-1], len(hits))
    for k, v in hits[:12]:
        print(k, v)
