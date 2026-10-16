# -*- coding: utf-8 -*-
"""DOCX 通用构建工具：中文排版、样式、表格、封面、页眉页脚。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x80, 0x80, 0x80)

HEI = "微软雅黑"
SONG = "宋体"


def _font(run, name=SONG, size=10.5, bold=False, color=DARK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), name)
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), color)
    tcPr.append(el)


def cell_text(cell, text, bold=False, size=9, color=DARK, align=None, name=SONG):
    cell.text = ""
    text = fix_quotes(str(text))
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align:
        p.alignment = align
    for i, line in enumerate(str(text).split("\n")):
        if i:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            if align:
                p.alignment = align
        _font(p.add_run(line), name=name, size=size, bold=bold, color=color)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), "true")
    trPr.append(el)


def add_table(doc, headers, rows, widths_cm, size=9, header_fill="1F3864", zebra=True,
              align_center_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    for i, h in enumerate(headers):
        cell_text(hdr.cells[i], h, bold=True, size=size, color=RGBColor(0xFF, 0xFF, 0xFF),
                  align=WD_ALIGN_PARAGRAPH.CENTER, name=HEI)
        shade(hdr.cells[i], header_fill)
    set_repeat_header(hdr)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            al = WD_ALIGN_PARAGRAPH.CENTER if i in align_center_cols else None
            cell_text(cells[i], v, size=size, align=al)
            if zebra and r_i % 2 == 1:
                shade(cells[i], "F2F5FA")
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    return t


def fix_quotes(text):
    """把中文正文里误写的 ASCII 双引号自动转成中文引号，避免排版与语法问题。"""
    out, pair = [], 0
    for ch in text:
        if ch == '"':
            out.append("\u201c" if pair % 2 == 0 else "\u201d")
            pair += 1
        else:
            out.append(ch)
    return "".join(out)


def h(doc, text, level=1, page_break=False):
    text = fix_quotes(text)
    if page_break and level == 1:
        doc.add_page_break()
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    size = {0: 20, 1: 15, 2: 12.5, 3: 11}[level]
    _font(run, name=HEI, size=size, bold=True, color=NAVY if level <= 2 else BLUE)
    return p


def _emit_runs(p, text, size=10.5, color=DARK, italic=False):
    """写入一个段落的内容，支持用 **文字** 标记加粗（跨行安全）。"""
    for line_i, line in enumerate(str(text).split("\n")):
        target = p if line_i == 0 else p.add_paragraph()
        if line_i:
            target.paragraph_format.space_after = Pt(2)
        for idx, seg in enumerate(re.split(r"\*\*(.+?)\*\*", line)):
            if not seg:
                continue
            _font(target.add_run(seg), size=size, bold=(idx % 2 == 1),
                  color=color, italic=italic)


def para(doc, text, size=10.5, indent=True, bold=False, color=DARK, italic=False, space_after=6):
    text = fix_quotes(text)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.35
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if bold:
        _font(p.add_run(text), size=size, bold=True, color=color, italic=italic)
    else:
        _emit_runs(p, text, size=size, color=color, italic=italic)
    return p


def bullet(doc, text, size=10.5, level=0):
    text = fix_quotes(text)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75 + level * 0.6)
    _emit_runs(p, text, size=size)
    return p


def numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    _font(p.add_run(text), size=size)
    return p


def add_code(doc, lines, size=8):
    """插入等宽字体代码块（单格表格 + 浅灰底）。"""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    _font(p.add_run("\n".join(lines)), name="Consolas", size=size)
    shade(c, "F7F7F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def note(doc, text, size=9):
    text = fix_quotes(text)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    _font(p.add_run(text), size=size, color=GREY, italic=True)
    return p


def cover(doc, title, subtitle, info_pairs, doc_kind="需求规格说明书"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    _font(p.add_run("齐河县职业中等专业学校"), name=HEI, size=16, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    _font(p.add_run(title), name=HEI, size=26, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run(subtitle), name=HEI, size=13, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    _font(p.add_run(doc_kind), name=HEI, size=14, bold=True, color=DARK)
    doc.add_paragraph()
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in info_pairs:
        cells = t.add_row().cells
        cell_text(cells[0], k, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, name=HEI)
        shade(cells[0], "DCE6F1")
        cell_text(cells[1], v, size=10)
        cells[0].width = Cm(3.4)
        cells[1].width = Cm(10.6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    _font(p.add_run("本文件含学校业务与管理信息，仅限项目组内部使用"), size=9, color=GREY)


def setup(doc, footer_text):
    st = doc.styles["Normal"]
    st.font.name = SONG
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.6)
        s.right_margin = Cm(2.4)
        # 页眉
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(hp.add_run("“智园卫士”智慧校园管理平台"), size=8.5, color=GREY)
        # 页脚页码
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        _font(run, size=9, color=GREY)
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        run._element.addnext(fld)
        _font(fp.add_run("  |  " + footer_text), size=8.5, color=GREY)


def save(doc, path):
    doc.save(path)
    print("已生成:", path)
