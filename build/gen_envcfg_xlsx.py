# -*- coding: utf-8 -*-
"""生成《环境配置进度与验收台账》xlsx —— 《电脑环境配置详细方案》的配套勾选工具"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"F:\职业中专\大赛\智慧园区创意\项目文档\10_环境配置进度与验收台账.xlsx"

NAVY, BLUE, LIGHT, GREY = "1F3864", "2E5C8A", "DCE6F1", "F2F2F7"
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr_style(ws, row=1, fill=NAVY):
    for c in ws[row]:
        if c.value is None:
            continue
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def body(ws, first=2, last=None, size=9):
    last = last or ws.max_row
    for r in ws.iter_rows(min_row=first, max_row=last, max_col=ws.max_column):
        for c in r:
            c.font = Font(name="微软雅黑", size=size)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name="微软雅黑", size=13, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1, value=sub).font = Font(name="微软雅黑", size=9, color="808080")
    ws.row_dimensions[2].height = 16


def sheet(wb, name, title_text, sub, headers, rows, col_widths, center_cols=(), status_col=None):
    ws = wb.create_sheet(name)
    title(ws, title_text, sub, len(headers))
    r = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=r, column=i, value=h)
    hdr_style(ws, r)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=r + i, column=j, value=v)
    body(ws, r + 1, r + len(rows))
    for rr in range(r + 1, r + len(rows) + 1):
        for cc in center_cols:
            ws.cell(row=rr, column=cc).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths(ws, col_widths)
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    if status_col:
        dv = DataValidation(type="list", formula1='"未开始,进行中,已完成,不适用"', allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(status_col)
        dv.add(f"{col}{r+1}:{col}{r+len(rows)}")
    return ws


wb = Workbook()
wb.remove(wb.active)

# ============================================================
# 01 机器分工与配置进度（总台账）
# ============================================================
PROGRESS = [
    ("台式机", "通用基线", "2.1.1", "登记机器信息并确认硬件条件", "内存 16GB，D 盘可用 ≥150GB", "", ""),
    ("台式机", "通用基线", "2.1.2", "开启 CPU 虚拟化", "VirtualizationFirmwareEnabled = True", "", ""),
    ("台式机", "通用基线", "2.1.3", "系统更新、时间、电源、长路径", "时区正确、不休眠、长路径=1", "", ""),
    ("台式机", "通用基线", "2.2.2", "winget 批量安装基础软件", "git/python/node/docker 均有版本", "", ""),
    ("台式机", "通用基线", "2.3.1", "UTF-8 与 Python 编码环境变量", "chcp=65001，中文不乱码", "", ""),
    ("台式机", "通用基线", "2.4.1", "配置 pip 清华源", "pip config list 显示清华源", "", ""),
    ("台式机", "通用基线", "2.4.2", "配置 npm 淘宝源", "registry=npmmirror", "", ""),
    ("台式机", "通用基线", "2.5.1", "创建 D:\\campus 目录结构", "目录齐全", "", ""),
    ("台式机", "通用基线", "2.5.2", "Git 全局配置与克隆仓库", "git remote -v 正常", "", ""),
    ("台式机", "通用基线", "2.6.1", "VS Code 与必备扩展", "扩展列表含 Python/Vue/Docker", "", ""),
    ("台式机", "WSL2", "3.1.1", "安装 WSL2 与 Ubuntu 24.04", "wsl -l -v 显示 Running / VERSION 2", "", ""),
    ("台式机", "WSL2", "3.1.2", "初始化 Ubuntu 基础环境", "ffmpeg / git / curl 可用", "", ""),
    ("台式机", "WSL2", "3.1.3", "配置 .wslconfig 资源上限", "free -h 接近 8GB", "", ""),
    ("台式机", "Docker", "3.2.1", "配置 Docker Desktop（WSL2/加速/资源）", "docker info 显示 Registry Mirrors", "", ""),
    ("台式机", "Docker", "3.2.2", "编写 docker-compose.dev.yml 与 .env", "git check-ignore 命中 .env", "", ""),
    ("台式机", "Docker", "3.2.3", "启动并校验 5 个依赖服务", "全部 running/healthy", "", ""),
    ("台式机", "后端", "3.3.1", "安装 uv 并创建后端虚拟环境", "which python 指向 .venv", "", ""),
    ("台式机", "后端", "3.3.2", "安装后端依赖（FastAPI 等）", "关键库可导入", "", ""),
    ("台式机", "后端", "3.3.3", "配置 .env 并启动首个接口", "/docs 与 /health 正常", "", ""),
    ("台式机", "后端", "3.3.4", "Alembic 初始化与首次迁移", "数据库可见业务表", "", ""),
    ("台式机", "AI 训练", "3.4.1", "安装 NVIDIA 驱动", "nvidia-smi 显示 RTX 3050", "", ""),
    ("台式机", "AI 训练", "3.4.2", "创建训练环境并安装 PyTorch(cu12x)", "torch.cuda.is_available() = True", "", ""),
    ("台式机", "AI 训练", "3.4.3", "安装 Ultralytics / ONNX 工具链", "yolo version 正常", "", ""),
    ("台式机", "AI 训练", "3.4.4", "准备数据集并跑通最小训练", "生成 best.pt", "", ""),
    ("台式机", "AI 导出", "3.5.1", "导出 ONNX（opset 17）", "best.onnx 存在", "", ""),
    ("台式机", "AI 导出", "3.5.2", "验证 ONNX 与 PyTorch 一致性", "偏差 ≤1e-3", "", ""),
    ("台式机", "前端", "3.6.1", "启用 pnpm 并创建 Vue3 工程", "pnpm dev 可访问 5173", "", ""),
    ("台式机", "前端", "3.6.2", "配置 /api 代理并联调", "5173/api/health 返回后端数据", "", ""),
    ("台式机", "设备模拟", "3.7.1", "搭建伪视频流（MediaMTX+FFmpeg）", "ffprobe 可读取 rtsp 流", "", ""),
    ("台式机", "设备模拟", "3.7.2", "搭建 MQTT 设备模拟器", "MQTTX 持续收到消息", "", ""),
    ("台式机", "设备模拟", "3.7.3", "准备门禁/车牌桩服务", "桩服务接口可调用", "", ""),
    ("台式机", "网关", "3.8.1", "本地 Nginx 反代（可选）", "8080 可访问且 /api 正常", "", ""),
    ("台式机", "监控", "3.9.1", "Prometheus + Grafana（可选）", "Grafana 显示后端指标", "", ""),
    ("笔记本1", "通用基线", "2.x", "完成第 2 章通用基线（精简）", "13 项校验通过", "", ""),
    ("笔记本1", "WSL2", "4.1.1", "安装 WSL2 并限制 Docker 资源", "vmmem ≤10GB", "", ""),
    ("笔记本1", "服务", "4.1.2", "拉起依赖服务与平台", "/health 正常、前端可访问", "", ""),
    ("笔记本1", "演示", "4.2.1", "构建离线演示环境", "断网后 8080 完整可用", "", ""),
    ("笔记本1", "演示", "4.2.2", "演示脚本与录屏备份", "录屏文件可播放", "", ""),
    ("笔记本2", "通用基线", "2.x", "完成第 2 章通用基线（精简）", "Git/Node/镜像源正常", "", ""),
    ("笔记本2", "客户端", "5.1.2", "配置指向开发机的访问入口", "可访问 8000/5173/18083/9001", "", ""),
    ("笔记本2", "测试", "5.2.1", "多角色并发与越权测试", "越权返回 403", "", ""),
    ("笔记本2", "设备", "5.2.2", "设备联调工具箱与现场准备", "工具可连设备", "", ""),
    ("阿里云", "创建", "6.1.1", "创建 ECS 并首次登录", "SSH 可登录，80/443 可达", "", ""),
    ("阿里云", "加固", "6.1.2", "系统加固与时间/编码配置", "禁用 root 直登，时区正确", "", ""),
    ("阿里云", "加固", "6.1.3", "安装 Docker 与加速", "docker compose version 正常", "", ""),
    ("阿里云", "部署", "6.2.1", "部署依赖服务与应用", "容器 healthy，/health 正常", "", ""),
    ("阿里云", "部署", "6.2.2", "配置 Nginx 与 HTTPS", "证书有效且自动续期可用", "", ""),
    ("阿里云", "部署", "6.2.3", "配置 MQTT over TLS 回传通道", "网关可 TLS 上报", "", ""),
    ("阿里云", "运维", "6.2.4", "配置自动备份与监控告警", "备份生成、告警可达", "", ""),
    ("机房", "母盘", "7.1.1", "母盘机最小集合环境", "可完成一次开发与提交", "", ""),
    ("机房", "离线包", "7.1.2", "制作离线包（wheels/pnpm/vsix）", "断网机器可装依赖", "", ""),
    ("机房", "脚本", "7.1.3", "批量安装与修复脚本", "脚本在母盘机跑通", "", ""),
    ("机房", "集中服务", "7.2.1", "集中开发服务器与账号隔离", "学生机可连且互相隔离", "", ""),
    ("现场", "服务器", "8.1.1", "系统安装与基础加固", "数据盘自动挂载、ufw 正确", "", ""),
    ("现场", "部署", "8.1.2", "部署平台与初始化数据", "基础数据导入完整", "", ""),
    ("现场", "网络", "8.2.1", "网络分区与端口放行", "设备互通、业务网不受影响", "", ""),
    ("现场", "设备", "8.2.2", "协议确认与联调准备", "每类设备协议明确", "", ""),
    ("现场", "运维", "8.3.1", "制作并演练回滚方案", "回滚 ≤15 分钟", "", ""),
    ("整体", "联调", "9.1", "端到端联调 10 个场景", "全部可演示通过", "", ""),
    ("整体", "自检", "9.2", "运行 check-env.py 自检", "通过率 100%", "", ""),
    ("整体", "归档", "9.3", "环境交付物归档", "9 类交付物齐全", "", ""),
]
sheet(wb, "01-配置进度总台账", "环境配置进度总台账",
      "版本 V1.0 ｜ 对应《电脑环境配置详细方案》章节 ｜ 状态列可下拉选择：未开始 / 进行中 / 已完成 / 不适用",
      ["机器", "阶段", "章节号", "配置项", "通过标准", "状态", "完成日期", "责任人", "备注"],
      [list(p) + ["", ""] for p in PROGRESS],
      [11, 11, 8, 40, 40, 10, 12, 10, 24], center_cols=(1, 2, 3, 6, 7, 8), status_col=6)

# ============================================================
# 02 台式机验收清单（25 项）
# ============================================================
DESKTOP = [
    ("D-01", "WSL2 正常", "wsl -l -v", "Ubuntu-24.04 Running，VERSION=2"),
    ("D-02", "WSL 资源限制生效", "wsl -d Ubuntu-24.04 -- free -h", "内存接近 .wslconfig 设定值"),
    ("D-03", "apt 源为国内", "grep -r aliyun /etc/apt/sources.list.d/", "有 aliyun 记录"),
    ("D-04", "Docker 可用", "docker run --rm hello-world", "输出 Hello from Docker!"),
    ("D-05", "Docker 加速生效", "docker info | Select-String 'Registry Mirrors'", "显示专属加速地址"),
    ("D-06", "PG16 + pgvector", "docker exec campus-pg psql -U campus -d campus -c '\\dx'", "扩展列表含 vector"),
    ("D-07", "Redis 7 连通", "docker exec campus-redis redis-cli -a <口令> ping", "PONG"),
    ("D-08", "EMQX 连通", "curl -s http://localhost:18083/status", "返回状态信息"),
    ("D-09", "MinIO 存活", "curl -s http://localhost:9000/minio/health/live", "HTTP 200"),
    ("D-10", "后端依赖齐全", "uv run python -c \"import fastapi, sqlalchemy, asyncpg\"", "无报错"),
    ("D-11", "后端服务可启", "curl -s http://localhost:8000/health", "status=ok"),
    ("D-12", "Swagger 可访问", "浏览器打开 /docs", "接口文档正常"),
    ("D-13", "数据库迁移成功", "uv run alembic current", "显示最新版本号"),
    ("D-14", "NVIDIA 驱动正常", "nvidia-smi", "显示 RTX 3050 与驱动版本"),
    ("D-15", "PyTorch CUDA 可用", "python -c \"import torch;print(torch.cuda.is_available())\"", "True"),
    ("D-16", "Ultralytics 可用", "yolo version", "输出版本号"),
    ("D-17", "最小训练跑通", "ls runs/detect/*/weights/best.pt", "权重文件存在"),
    ("D-18", "ONNX 导出成功", "ls runs/detect/*/weights/best.onnx", "文件存在"),
    ("D-19", "推理一致性达标", "对比脚本输出", "偏差 ≤1e-3"),
    ("D-20", "前端可启动", "浏览器访问 http://localhost:5173", "页面正常"),
    ("D-21", "前后端联调通", "curl -s http://localhost:5173/api/health", "返回后端 JSON"),
    ("D-22", "伪视频可拉流", "ffprobe rtsp://localhost:8554/cam1", "输出视频流信息"),
    ("D-23", "MQTT 模拟可上报", "MQTTX 订阅 iot/#", "持续收到消息"),
    ("D-24", "桩服务可访问", "curl -X POST http://localhost:9100/api/device/access/report", "返回模拟数据"),
    ("D-25", "环境自检通过", "python 环境工具\\scripts\\check-env.py --role dev", "通过率 100%"),
]
sheet(wb, "02-台式机验收", "台式机（主力开发机）验收清单",
      "对应方案第 3.10 节 ｜ 25 项全部通过方可进入功能开发",
      ["编号", "验收项", "校验命令", "通过标准", "结果", "实际输出/备注", "验收人", "日期"],
      [list(d) + ["", "", "", ""] for d in DESKTOP],
      [8, 22, 46, 34, 8, 30, 10, 12], center_cols=(1, 5, 7, 8), status_col=None)

# ============================================================
# 03 其他机器验收清单
# ============================================================
OTHERS = [
    ("笔记本1", "S-01", "通用基线 13 项全过", "见第 2.7 节"),
    ("笔记本1", "S-02", "WSL2 + Ubuntu 正常", "wsl -l -v 显示 VERSION 2"),
    ("笔记本1", "S-03", "Docker 内存限制生效", "vmmem 占用 ≤10GB"),
    ("笔记本1", "S-04", "依赖服务全起", "docker compose ps 全部 healthy"),
    ("笔记本1", "S-05", "后端 /health 正常", "status=ok"),
    ("笔记本1", "S-06", "前端可访问", "5173 或 8080 页面正常"),
    ("笔记本1", "S-07", "离线演示可用", "断网后 8080 完整可用"),
    ("笔记本1", "S-08", "镜像已本地备份", "images-demo.tar 可导入"),
    ("笔记本1", "S-09", "演示录屏已备份", "录屏文件可播放"),
    ("笔记本1", "S-10", "环境自检通过", "check-env.py --role server ≥90%"),
    ("笔记本2", "C-01", "通用基线（精简）通过", "Git/Node/镜像源正常"),
    ("笔记本2", "C-02", "能访问开发机服务", "8000/5173/18083/9001 可打开"),
    ("笔记本2", "C-03", "多角色并发测试可用", "5 类角色可同时登录"),
    ("笔记本2", "C-04", "越权测试无异常", "跨角色访问返回 403"),
    ("笔记本2", "C-05", "MQTT 工具可连", "MQTTX 能订阅消息"),
    ("笔记本2", "C-06", "串口工具可用", "能打开串口设备"),
    ("笔记本2", "C-07", "现场工具齐备", "网线/USB-485/热点已随身"),
    ("阿里云", "K-01", "SSH 安全加固完成", "禁用 root 与口令登录"),
    ("阿里云", "K-02", "时区与时间同步", "Asia/Shanghai 且 NTP 正常"),
    ("阿里云", "K-03", "安全组最小开放", "仅 22/80/443（8883 限来源）"),
    ("阿里云", "K-04", "Docker 与 Compose 可用", "版本命令正常"),
    ("阿里云", "K-05", "依赖服务全起", "全部 healthy"),
    ("阿里云", "K-06", "后端接口可访问", "/health 返回 ok"),
    ("阿里云", "K-07", "HTTPS 有效", "无证书警告且可自动续期"),
    ("阿里云", "K-08", "MQTT over TLS 可用", "校内网关可 TLS 上报"),
    ("阿里云", "K-09", "自动备份生效", "备份文件生成且保留策略正确"),
    ("阿里云", "K-10", "恢复演练通过", "恢复数据与源库一致"),
    ("阿里云", "K-11", "监控告警可达", "手机收到阈值告警"),
    ("阿里云", "K-12", "日志与磁盘", "日志按天切割，磁盘告警可用"),
    ("机房", "L-01", "母盘环境完成", "可独立完成一次开发与提交"),
    ("机房", "L-02", "批量部署完成", "抽查 5 台环境一致"),
    ("机房", "L-03", "离线包可用", "断网可安装依赖"),
    ("机房", "L-04", "集中开发服务器可用", "账号隔离，DB/API 可连"),
    ("机房", "L-05", "修复脚本可用", "5 分钟内恢复环境"),
    ("机房", "L-06", "作业不丢失", "Git 提交，还原不影响成果"),
    ("机房", "L-07", "学生账号隔离", "A 无法访问 B 的数据"),
    ("现场", "F-01", "平台服务全起", "容器 healthy，重启自恢复"),
    ("现场", "F-02", "数据库与基础数据", "组织/场所/人员/台账完整"),
    ("现场", "F-03", "门禁联调", "刷卡刷脸 P95 ≤300ms"),
    ("现场", "F-04", "断网离线放行", "断网 30 分钟可放行且补传完整"),
    ("现场", "F-05", "视频接入", "预览延时 ≤1.5 秒"),
    ("现场", "F-06", "AI 算法达标", "准确率与误报率符合要求"),
    ("现场", "F-07", "通知通道", "短信/微信送达可查回执"),
    ("现场", "F-08", "定时任务", "归寝判定/报表/备份按计划执行"),
    ("现场", "F-09", "备份与回滚演练", "RPO ≤15min、RTO ≤2h、回滚 ≤15min"),
    ("现场", "F-10", "培训与文档移交", "三类培训签到，手册齐全"),
]
sheet(wb, "03-其他机器验收", "笔记本 1 / 笔记本 2 / 云 / 机房 / 现场 验收清单",
      "对应方案第 4.3、5.3、6.3、7.3、8.4 节",
      ["环境", "编号", "验收项", "通过标准", "结果", "实际输出/备注", "验收人", "日期"],
      [list(o) + ["", "", "", ""] for o in OTHERS],
      [10, 8, 26, 40, 8, 30, 10, 12], center_cols=(1, 2, 5, 7, 8))

# ============================================================
# 04 端到端联调场景
# ============================================================
E2E = [
    ("E-01", "人员建档与门禁通行", "新增测试人员并授权 → 模拟设备上报该人通行", "生成通行记录，在校状态变为在校"),
    ("E-02", "请假与离校拦截", "请假审批通过 → 上报离校通行；再测试未请假人员离校", "请假者放行；未请假者被拦截至事件"),
    ("E-03", "通行即考勤", "上报教职工上下班通行记录", "考勤自动生成签到签退"),
    ("E-04", "AI 事件闭环", "向伪视频流播放烟火视频 → 算法识别", "事件生成、派单、处置闭环、证据留档"),
    ("E-05", "设备心跳与离线告警", "停止模拟设备上报", "阈值内产生设备离线告警"),
    ("E-06", "通知链路", "触发未归预警规则", "站内即时到达；短信/微信送达可查回执"),
    ("E-07", "权限与数据隔离", "用学生账号访问管理接口", "返回 403 且审计日志有记录"),
    ("E-08", "断网可用性", "断开平台与门禁网络 30 分钟", "门禁离线放行，恢复后补传完整"),
    ("E-09", "备份与恢复", "对云端或现场库做恢复演练", "数据一致且满足 RTO"),
    ("E-10", "性能抽测", "对门禁鉴权接口施加并发压力", "P95 ≤300ms，无失败请求"),
]
sheet(wb, "04-端到端联调", "端到端联调场景验收（10 个场景）",
      "对应方案第 9.1 节 ｜ 每条场景必须现场演示通过并留证据",
      ["编号", "联调场景", "操作步骤", "通过标准", "结果", "证据（截图/日志路径）", "执行人", "日期"],
      [list(e) + ["", "", "", ""] for e in E2E],
      [8, 22, 44, 38, 8, 34, 10, 12], center_cols=(1, 5, 7, 8))

# ============================================================
# 05 软件安装清单
# ============================================================
SOFT = [
    (1, "基础", "Git for Windows", "最新稳定版", "全部机器", "版本管理", "winget install --id Git.Git -e", "git --version"),
    (2, "基础", "VS Code", "最新稳定版", "全部机器", "主开发工具", "winget install --id Microsoft.VisualStudioCode -e", "code --version"),
    (3, "基础", "PowerShell 7", "最新版", "全部机器", "命令行环境", "winget install --id Microsoft.PowerShell -e", "$PSVersionTable"),
    (4, "基础", "Windows Terminal", "最新版", "全部机器", "终端体验", "winget install --id Microsoft.WindowsTerminal -e", "wt -v"),
    (5, "基础", "Python 3.12", "3.12.x", "台式机/笔记本1/机房", "后端与脚本", "winget install --id Python.Python.3.12 -e", "python --version"),
    (6, "基础", "Python 3.11", "3.11.x", "台式机", "AI 训练生态兼容", "winget install --id Python.Python.3.11 -e", "py -3.11 --version"),
    (7, "基础", "Node.js LTS", "22.x", "全部机器", "前端构建", "winget install --id OpenJS.NodeJS.LTS -e", "node -v"),
    (8, "基础", "pnpm", "最新版", "台式机/笔记本1/机房", "前端包管理", "corepack prepare pnpm@latest --activate", "pnpm -v"),
    (9, "基础", "uv", "最新版", "台式机/笔记本1", "Python 依赖管理", "irm https://astral.sh/uv/install.ps1 | iex", "uv --version"),
    (10, "容器", "WSL2 + Ubuntu 24.04", "24.04 LTS", "台式机/笔记本1", "Linux 开发底座", "wsl --install -d Ubuntu-24.04", "wsl -l -v"),
    (11, "容器", "Docker Desktop", "最新稳定版", "台式机/笔记本1", "容器运行", "winget install --id Docker.DockerDesktop -e", "docker version"),
    (12, "数据", "DBeaver Community", "最新版", "全部机器", "数据库客户端", "winget install --id dbeaver.dbeaver -e", "启动并连接"),
    (13, "数据", "PostgreSQL 16 + pgvector", "pgvector/pgvector:pg16", "台式机/笔记本1/云/现场", "业务主库与向量检索", "docker compose up -d postgres", "docker exec campus-pg psql -c '\\dx'"),
    (14, "数据", "Redis 7", "redis:7-alpine", "台式机/笔记本1/云/现场", "缓存与白名单", "docker compose up -d redis", "redis-cli ping"),
    (15, "数据", "TDengine", "tdengine/tdengine", "台式机/现场", "时序数据（V2 启用）", "docker compose up -d tdengine", "curl 6041 探活"),
    (16, "消息", "EMQX 5", "emqx/emqx:5", "台式机/笔记本1/云/现场", "MQTT 设备接入", "docker compose up -d emqx", "curl 18083/status"),
    (17, "存储", "MinIO", "minio/minio", "台式机/笔记本1/云/现场", "对象存储", "docker compose up -d minio", "curl 9000/minio/health/live"),
    (18, "网关", "Nginx", "nginx:alpine", "笔记本1/云/现场", "反向代理与静态", "docker run nginx:alpine", "nginx -t"),
    (19, "AI", "NVIDIA 驱动", "官方最新", "台式机", "CUDA 支持", "NVIDIA 官网下载安装", "nvidia-smi"),
    (20, "AI", "PyTorch (cu12x)", "2.x", "台式机", "模型训练", "见 pytorch.org 生成的命令", "torch.cuda.is_available()"),
    (21, "AI", "Ultralytics", "最新版", "台式机", "YOLO 训练与导出", "uv pip install ultralytics", "yolo version"),
    (22, "AI", "ONNX Runtime", "最新版", "台式机", "边缘推理", "uv pip install onnxruntime", "import onnxruntime"),
    (23, "AI", "LabelImg / X-AnyLabeling", "最新版", "台式机", "数据标注", "uv pip install labelImg", "启动标注工具"),
    (24, "AI", "OpenCV / NumPy / Pillow", "最新稳定版", "台式机", "图像处理", "uv pip install opencv-python numpy pillow", "import cv2"),
    (25, "前端", "Vite + Vue3 + TS", "最新版", "全部机器", "前端工程", "pnpm create vite web-admin --template vue-ts", "pnpm dev"),
    (26, "前端", "Element Plus / Pinia / Vue Router / Axios", "最新版", "全部机器", "UI 与状态", "pnpm add element-plus pinia vue-router axios", "页面正常"),
    (27, "前端", "ECharts", "5.x", "全部机器", "图表与大屏", "pnpm add echarts", "图表渲染"),
    (28, "前端", "uni-app", "最新版", "按需", "小程序端", "HBuilderX 或 CLI", "小程序预览"),
    (29, "设备", "MQTTX", "最新版", "台式机/笔记本1/2", "MQTT 调试", "官网下载", "连接并收发"),
    (30, "设备", "FFmpeg", "最新稳定版", "台式机/笔记本1", "伪视频与视频处理", "apt install ffmpeg / 官网", "ffmpeg -version"),
    (31, "设备", "MediaMTX", "最新版", "台式机/笔记本1", "RTSP 服务器", "docker run bluenviron/mediamtx", "ffprobe rtsp://"),
    (32, "设备", "ONVIF Device Manager", "最新版", "笔记本2", "摄像机发现", "GitHub 下载", "扫描到设备"),
    (33, "调试", "Apifox / Postman", "最新版", "全部机器", "接口调试", "官网下载", "调用接口"),
    (34, "调试", "MobaXterm / WinSCP", "最新版", "全部机器", "SSH 与文件传输", "官网下载", "连接服务器"),
    (35, "调试", "SSCOM / XCOM", "最新版", "笔记本2", "串口调试", "网络下载", "打开串口"),
    (36, "调试", "Wireshark", "最新版", "台式机/笔记本2", "协议抓包", "官网下载", "抓到报文"),
    (37, "演示", "OBS Studio", "最新版", "台式机/笔记本1", "演示录屏", "winget install --id OBSProject.OBSStudio -e", "录制成功"),
    (38, "云侧", "Certbot", "最新版", "云服务器", "HTTPS 证书", "apt install certbot python3-certbot-nginx", "certbot renew --dry-run"),
    (39, "云侧", "云监控与快照", "控制台配置", "云服务器", "可用性与数据保护", "阿里云控制台", "收到告警"),
    (40, "办公", "Office / WPS / Markdown 编辑器", "最新版", "全部机器", "文档编写", "官网下载", "文档可编辑"),
]
sheet(wb, "05-软件安装清单", "软件安装与校验清单（40 项）",
      "对应方案第 2.2、3.x、4.x、5.x 节 ｜ 状态列可下拉选择",
      ["序号", "类别", "软件/组件", "版本要求", "适用机器", "用途", "安装命令/方式", "校验命令", "状态"],
      [list(s) + [""] for s in SOFT],
      [6, 8, 30, 20, 22, 22, 44, 30, 10], center_cols=(1, 2, 9), status_col=9)

# ============================================================
# 06 命令速查
# ============================================================
CMD = [
    ("系统检查", "检查虚拟化是否开启", "(Get-CimInstance Win32_Processor).VirtualizationFirmwareEnabled"),
    ("系统检查", "查看磁盘可用空间", "Get-PSDrive -PSProvider FileSystem | Select Name,@{n='可用GB';e={[math]::Round($_.Free/1GB,1)}}"),
    ("系统检查", "确认 UTF-8 代码页", "chcp"),
    ("WSL", "查看发行版与版本", "wsl -l -v"),
    ("WSL", "进入 Ubuntu", "wsl -d Ubuntu-24.04"),
    ("WSL", "限制资源后重启 WSL", "wsl --shutdown"),
    ("镜像源", "查看 pip 源", "pip config list"),
    ("镜像源", "查看 npm 源", "npm config get registry"),
    ("Docker", "查看 Docker 与加速配置", "docker info | Select-String 'Server Version','Registry Mirrors'"),
    ("Docker", "启动开发依赖服务", "docker compose -f docker-compose.dev.yml up -d"),
    ("Docker", "查看容器状态", "docker compose -f docker-compose.dev.yml ps"),
    ("Docker", "查看服务日志", "docker logs -f campus-pg"),
    ("Docker", "停止并保留数据", "docker compose -f docker-compose.dev.yml down"),
    ("数据库", "进入 PostgreSQL", "docker exec -it campus-pg psql -U campus -d campus"),
    ("数据库", "查看扩展（pgvector）", "docker exec campus-pg psql -U campus -d campus -c '\\dx'"),
    ("数据库", "备份数据库", "docker exec campus-pg pg_dump -U campus -d campus | gzip > backup.sql.gz"),
    ("数据库", "恢复数据库", "gunzip -c backup.sql.gz | docker exec -i campus-pg psql -U campus -d campus"),
    ("缓存", "Redis 连通性测试", "docker exec campus-redis redis-cli -a <口令> ping"),
    ("消息", "EMQX 状态", "curl -s http://localhost:18083/status"),
    ("存储", "MinIO 健康检查", "curl -s http://localhost:9000/minio/health/live"),
    ("后端", "启动开发服务器", "uv run uvicorn app.main:app --reload --port 8000"),
    ("后端", "执行数据库迁移", "uv run alembic upgrade head"),
    ("后端", "健康检查", "curl -s http://localhost:8000/health"),
    ("后端", "查看接口文档", "浏览器打开 http://localhost:8000/docs"),
    ("前端", "启动开发服务器", "pnpm dev --host"),
    ("前端", "构建生产产物", "pnpm build"),
    ("AI", "验证 GPU 与 CUDA", "python -c \"import torch;print(torch.cuda.is_available(), torch.cuda.get_device_name(0))\""),
    ("AI", "最小训练（10 轮）", "yolo detect train model=yolo11n.pt data=data.yaml epochs=10 imgsz=640 batch=8 device=0 amp=True"),
    ("AI", "导出 ONNX", "yolo export model=best.pt format=onnx opset=17 dynamic=True simplify=True"),
    ("AI", "查看显存占用", "nvidia-smi --query-gpu=memory.used,memory.total --format=csv"),
    ("设备模拟", "推伪视频流", "ffmpeg -re -stream_loop -1 -i sample.mp4 -c copy -f rtsp rtsp://localhost:8554/cam1"),
    ("设备模拟", "探测 RTSP 流", "ffprobe -rtsp_transport tcp -i rtsp://localhost:8554/cam1 -show_streams"),
    ("设备模拟", "订阅 MQTT 主题", "docker run --rm eclipse-mosquitto mosquitto_sub -h <IP> -t 'iot/#' -v"),
    ("自检", "环境自检（开发机）", "python 环境工具\\scripts\\check-env.py --role dev"),
    ("自检", "环境自检（输出 JSON）", "python 环境工具\\scripts\\check-env.py --json > env.json"),
    ("自检", "一键基础安装（管理员）", ".\\环境工具\\scripts\\install-base.ps1 -Role dev"),
]
sheet(wb, "06-命令速查", "常用命令速查表",
      "按场景分类，可直接复制执行；<口令>/<IP> 等占位符请替换为实际值",
      ["场景", "用途", "命令"],
      [list(c) for c in CMD],
      [12, 34, 90], center_cols=(1,))

# ============================================================
# 07 服务口令与端口台账
# ============================================================
SECRET = [
    ("PostgreSQL", "campus-pg", "5432", "campus / campus", "写入 deploy/.env 的 PG_PASSWORD", "仅内网"),
    ("Redis", "campus-redis", "6379", "无用户名", "写入 .env 的 REDIS_PASSWORD，开启 AOF", "仅内网"),
    ("EMQX Dashboard", "campus-emqx", "18083", "admin", "首次登录后立即修改；写入 .env", "仅内网/云限来源"),
    ("EMQX MQTT", "campus-emqx", "1883 / 8883", "设备级账号", "为每台网关分配独立账号，ACL 限制主题", "1883 仅内网；8883 云上限来源"),
    ("MinIO", "campus-minio", "9000 / 9001", "自定义强密钥", "桶权限设为私有；写入 .env", "仅内网"),
    ("TDengine", "campus-tdengine", "6030 / 6041", "root", "启用后立即修改默认口令", "仅内网"),
    ("后端 JWT", "backend", "8000", "无", "JWT_SECRET 用 32 位以上随机串", "应用层鉴权"),
    ("Ubuntu 运维账号", "各服务器", "22", "deploy（sudo）", "仅密钥登录，禁用 root 直登", "限来源 IP"),
    ("微信小程序", "云", "443", "AppID/Secret", "Secret 只存服务端环境变量", "不写入前端"),
    ("阿里云 AccessKey", "云", "—", "RAM 子账号", "按最小权限授权，用于备份上传 OSS", "不写入代码"),
]
sheet(wb, "07-口令与端口台账", "服务口令、端口与开放范围台账",
      "本表用于登记配置信息（不含真实口令）｜ 真实口令请单独加密保存，禁止写入本表或提交到代码仓库",
      ["服务", "容器/主机名", "端口", "账号", "口令管理方式", "开放范围", "实际配置日期", "备注"],
      [list(s) + ["", ""] for s in SECRET],
      [16, 18, 14, 16, 46, 26, 14, 26], center_cols=(3, 7))

# ============================================================
# 08 IP 地址台账
# ============================================================
IPROWS = [
    ("摄像机", "教学楼一层东走廊", "", "", "", "", "RTSP 554"),
    ("摄像机", "宿舍楼出入口", "", "", "", "", "RTSP 554"),
    ("人脸门禁一体机", "校门人行通道 1", "", "", "", "", "HTTP 80 / 韦根"),
    ("车牌识别一体机", "校门车行道", "", "", "", "", "HTTP 80"),
    ("道闸控制器", "校门车行道", "", "", "", "", "TCP"),
    ("NVR", "机房", "", "", "", "", "HTTP 80 / RTSP"),
    ("边缘 AI 盒子", "弱电间 A", "", "", "", "", "HTTP / MQTT"),
    ("环境传感器网关", "教学楼", "", "", "", "", "Modbus TCP 502"),
    ("智能电表网关", "配电间", "", "", "", "", "Modbus TCP 502"),
    ("现场服务器", "中心机房", "", "", "", "", "22 / 443 / 8000"),
    ("", "", "", "", "", "", ""),
    ("", "", "", "", "", "", ""),
    ("", "", "", "", "", "", ""),
    ("", "", "", "", "", "", ""),
    ("", "", "", "", "", "", ""),
]
sheet(wb, "08-IP地址台账", "设备与主机 IP 地址台账（现场勘察时填写）",
      "编号—位置—IP—MAC 必须登记完整，否则后期排查离线设备将非常困难",
      ["设备类型", "安装位置", "设备编号", "IP 地址", "MAC", "所属 VLAN", "开放端口", "责任人", "登记日期"],
      [list(i) + ["", ""] for i in IPROWS],
      [18, 24, 14, 16, 20, 12, 18, 10, 12], center_cols=(3, 4, 5, 6, 8, 9))

wb.save(OUT)
print("已生成:", OUT)
print("工作表:", wb.sheetnames)
