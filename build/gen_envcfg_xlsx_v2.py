# -*- coding: utf-8 -*-
"""生成 10_环境配置进度与验收台账（单机版）.xlsx —— 与文档 09 V2 的章节编号同步"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"F:\职业中专\大赛\智慧园区创意\项目文档\10_环境配置进度与验收台账（单机版）.xlsx"

NAVY, BLUE, LIGHT = "1F3864", "2E5C8A", "DCE6F1"
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr_style(ws, row=1, fill=NAVY):
    for c in ws[row]:
        if c.value is None:
            continue
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def body(ws, first=2, last=None, size=9):
    last = last or ws.max_row
    for r in ws.iter_rows(min_row=first, max_row=last, max_col=ws.max_column):
        for c in r:
            c.font = Font(name="微软雅黑", size=size)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name="微软雅黑", size=13, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1, value=sub).font = Font(name="微软雅黑", size=9, color="808080")


def sheet(wb, name, t, sub, headers, rows, col_widths, center_cols=(), status_col=None):
    ws = wb.create_sheet(name)
    title(ws, t, sub, len(headers))
    r = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=r, column=i, value=h)
    hdr_style(ws, r)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=r + i, column=j, value=v)
    body(ws, r + 1, r + len(rows))
    for rr in range(r + 1, r + len(rows) + 1):
        for cc in center_cols:
            ws.cell(row=rr, column=cc).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths(ws, col_widths)
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    if status_col:
        dv = DataValidation(type="list", formula1='"未开始,进行中,已完成,不适用"', allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(status_col)
        dv.add(f"{col}{r+1}:{col}{r+len(rows)}")
    return ws


wb = Workbook()
wb.remove(wb.active)

# ============================================================
# 00 使用说明与新旧章节对照
# ============================================================
CROSS = [
    ("通用基线（软件、编码、镜像源、目录、Git）", "09 V1 第 2 章", "09 V2 第 2 章", "编号未变"),
    ("WSL2 + Ubuntu 24.04", "09 V1 第 3.1 节", "09 V2 第 4.1 节", "并入台式机全栈主线"),
    ("Docker Desktop 与依赖服务全栈", "09 V1 第 3.2 节", "09 V2 第 4.2 节", "并入台式机全栈主线"),
    ("后端 Python 环境", "09 V1 第 3.3 节", "09 V2 第 4.3 节", "后端与前端合并一节"),
    ("AI 训练环境与 ONNX 导出", "09 V1 第 3.4—3.5 节", "09 V2 第 4.4 节", "合并"),
    ("前端开发环境", "09 V1 第 3.6 节", "09 V2 第 4.3 节", "与后端合并"),
    ("设备模拟与本地网关", "09 V1 第 3.7—3.9 节", "09 V2 第 4.5 节", "合并"),
    ("台式机验收（25 项）", "09 V1 第 3.10 节", "09 V2 第 9 章", "位置后移"),
    ("单机快速验收（15 项）", "—", "09 V2 第 4.6 节", "新增，半天可完成"),
    ("笔记本 1（演示机）", "09 V1 第 4 章", "09 V2 第 5 章", "改为“外援环境一”，可选"),
    ("笔记本 2（客户端）", "09 V1 第 5 章", "09 V2 第 6 章", "改为“外援环境二：协作环境”，可选"),
    ("阿里云 ECS（服务器）", "09 V1 第 6 章", "09 V2 第 7 章", "上线阶段才需要"),
    ("学生机房批量部署", "09 V1 第 7 章", "09 V2 第 8 章", "按需"),
    ("整体联调与自检", "09 V1 第 9 章", "09 V2 第 9 章", "编号未变"),
    ("FAQ 与附录", "09 V1 附录 A—D", "09 V2 附录 A—D", "编号未变"),
    ("三类环境总览与快速开始", "—", "文档 11", "新增决策指南"),
    ("协作同学环境", "—", "文档 12", "新增，可整份转发"),
]
sheet(wb, "00-使用说明与章节对照", "环境配置进度与验收台账（单机全栈版）",
      "版本 V2.0 ｜ 与《电脑环境配置详细方案》V2.0 的章节编号一一对应 ｜ 若你手上的旧文档还是 V1 编号，请先看本表对照",
      ["配置内容", "旧编号（09 V1）", "新编号（09 V2）", "变化说明"],
      [list(c) for c in CROSS],
      [44, 22, 22, 26], center_cols=(2, 3))

# ============================================================
# 01 配置进度总台账（单机主线优先）
# ============================================================
PROGRESS = [
    ("台式机（主线）", "通用基线", "2.1.1", "登记机器信息并确认硬件条件", "内存 16GB，D 盘可用 ≥150GB", "", "", ""),
    ("台式机（主线）", "通用基线", "2.1.2", "开启 CPU 虚拟化", "VirtualizationFirmwareEnabled = True", "", "", ""),
    ("台式机（主线）", "通用基线", "2.1.3", "系统更新、时间、电源、长路径", "时区正确、不休眠、长路径=1", "", "", ""),
    ("台式机（主线）", "通用基线", "2.2.2", "winget 批量安装基础软件", "git/python/node/docker 均有版本", "", "", ""),
    ("台式机（主线）", "通用基线", "2.3.1", "UTF-8 与 Python 编码环境变量", "chcp=65001，中文不乱码", "", "", ""),
    ("台式机（主线）", "通用基线", "2.4.1", "配置 pip 清华源", "pip config list 显示清华源", "", "", ""),
    ("台式机（主线）", "通用基线", "2.4.2", "配置 npm 淘宝源", "registry=npmmirror", "", "", ""),
    ("台式机（主线）", "通用基线", "2.5.1", "创建 D:\\campus 目录结构", "目录齐全", "", "", ""),
    ("台式机（主线）", "通用基线", "2.5.2", "Git 全局配置与克隆仓库", "git remote -v 正常", "", "", ""),
    ("台式机（主线）", "通用基线", "2.6.1", "VS Code 与必备扩展", "扩展含 Python/Vue/Docker", "", "", ""),
    ("台式机（主线）", "单机全栈 4.1", "4.1.1", "安装 WSL2 与 Ubuntu 24.04", "wsl -l -v 显示 Running / VERSION 2", "", "", ""),
    ("台式机（主线）", "单机全栈 4.1", "4.1.2", "初始化 Ubuntu 基础环境", "ffmpeg / git / curl 可用", "", "", ""),
    ("台式机（主线）", "单机全栈 4.1", "4.1.3", "配置 .wslconfig 与代码位置", "free -h 接近 8GB，代码在 ~/projects", "", "", ""),
    ("台式机（主线）", "单机全栈 4.2", "4.2.1", "配置 Docker Desktop（WSL2/加速/资源）", "docker info 显示 Registry Mirrors", "", "", ""),
    ("台式机（主线）", "单机全栈 4.2", "4.2.2", "编写 docker-compose.dev.yml 与 .env", "git check-ignore 命中 .env", "", "", ""),
    ("台式机（主线）", "单机全栈 4.2", "4.2.3", "启动并校验依赖服务", "PG/Redis/EMQX/MinIO 全部 healthy", "", "", ""),
    ("台式机（主线）", "单机全栈 4.3", "4.3.1", "安装 uv 并创建后端虚拟环境", "which python 指向 .venv", "", "", ""),
    ("台式机（主线）", "单机全栈 4.3", "4.3.2", "安装后端依赖（FastAPI 等）", "关键库可导入", "", "", ""),
    ("台式机（主线）", "单机全栈 4.3", "4.3.3", "配置 .env 并启动接口", "/docs 与 /health 正常", "", "", ""),
    ("台式机（主线）", "单机全栈 4.3", "4.3.4", "Alembic 初始化与首次迁移", "数据库可见业务表", "", "", ""),
    ("台式机（主线）", "单机全栈 4.3", "4.3.5", "启用 pnpm 并创建前端工程", "5173 可访问", "", "", ""),
    ("台式机（主线）", "单机全栈 4.3", "4.3.6", "配置 /api 代理并联调", "5173/api/health 返回后端数据", "", "", ""),
    ("台式机（主线）", "单机全栈 4.4", "4.4.1", "安装 NVIDIA 驱动", "nvidia-smi 显示 RTX 3050", "", "", ""),
    ("台式机（主线）", "单机全栈 4.4", "4.4.2", "创建训练环境并装 PyTorch(cu12x)", "torch.cuda.is_available() = True", "", "", ""),
    ("台式机（主线）", "单机全栈 4.4", "4.4.3", "安装 Ultralytics / ONNX 工具链", "yolo version 正常", "", "", ""),
    ("台式机（主线）", "单机全栈 4.4", "4.4.4", "准备数据集并跑通最小训练", "生成 best.pt", "", "", ""),
    ("台式机（主线）", "单机全栈 4.4", "4.4.5", "导出 ONNX 并验证一致性", "偏差 ≤1e-3", "", "", ""),
    ("台式机（主线）", "单机全栈 4.5", "4.5.1", "搭建伪视频流（MediaMTX+FFmpeg）", "ffprobe 可读取 rtsp 流", "", "", ""),
    ("台式机（主线）", "单机全栈 4.5", "4.5.2", "搭建 MQTT 模拟器与桩服务", "MQTTX 持续收到消息", "", "", ""),
    ("台式机（主线）", "单机全栈 4.5", "4.5.3", "本地 Nginx 与监控（可选）", "8080 可访问、Grafana 有指标", "", "", ""),
    ("台式机（主线）", "单机验收", "4.6", "15 项单机快速验收", "全部通过", "", "", ""),
    ("台式机（主线）", "全栈验收", "9.x", "25 项全栈验收 + 环境自检", "check-env.py 通过率 100%", "", "", ""),
    ("演示机（可选）", "演示环境", "5.1.1", "通用基线（精简）+ 依赖服务", "依赖服务 healthy", "", "", ""),
    ("演示机（可选）", "演示环境", "5.1.2", "构建前端产物 + Nginx 托管", "8080 可访问且 /api 正常", "", "", ""),
    ("演示机（可选）", "演示环境", "5.1.3", "断网演示包与录屏备份", "拔网线可演示、录屏可播放", "", "", ""),
    ("协作同学（可选）", "协作环境", "6.1.1", "安装基础工具与镜像源", "git/node/pnpm 正常", "", "", ""),
    ("协作同学（可选）", "协作环境", "6.1.2", "拉取代码并连开发机后端", "页面能加载后端数据", "", "", ""),
    ("协作同学（可选）", "协作环境", "6.1.3", "交接访问清单与协作约定", "对方能独立完成一次提交", "", "", ""),
    ("服务器（可选）", "云服务器", "7.1.1", "创建 ECS 并首次登录", "SSH 可登录，80/443 可达", "", "", ""),
    ("服务器（可选）", "云服务器", "7.1.2", "系统加固与时间/编码", "禁用 root 直登，时区正确", "", "", ""),
    ("服务器（可选）", "云服务器", "7.1.3", "安装 Docker 与 Compose", "compose version 正常", "", "", ""),
    ("服务器（可选）", "云服务器", "7.2.1", "部署依赖服务与应用", "容器 healthy，/health 正常", "", "", ""),
    ("服务器（可选）", "云服务器", "7.2.2", "配置 Nginx 与 HTTPS", "证书有效且可续期", "", "", ""),
    ("服务器（可选）", "云服务器", "7.2.3", "配置 MQTT over TLS 回传通道", "网关可 TLS 上报", "", "", ""),
    ("服务器（可选）", "云服务器", "7.2.4", "配置自动备份与监控告警", "备份生成、告警可达", "", "", ""),
    ("机房（按需）", "机房", "8.x", "母盘环境与批量部署", "抽查一致、可还原", "", "", ""),
    ("整体", "联调与归档", "9.x", "端到端联调与交付物归档", "场景全通过、9 类交付物齐全", "", "", ""),
]
sheet(wb, "01-配置进度总台账", "环境配置进度总台账（单机主线优先）",
      "状态列可下拉：未开始 / 进行中 / 已完成 / 不适用 ｜ 主线的 32 项完成即可开始开发，其余为可选",
      ["机器/环境", "阶段", "章节号", "配置项", "通过标准", "状态", "完成日期", "责任人", "备注"],
      [list(p) for p in PROGRESS],
      [16, 14, 9, 42, 38, 10, 12, 10, 22], center_cols=(2, 3, 6, 7, 8), status_col=6)

# ============================================================
# 02 单机快速验收（15 项）
# ============================================================
QUICK = [
    ("Q-01", "基础软件齐全", "git --version; python --version; node -v", "均有版本输出"),
    ("Q-02", "镜像源已配置", "pip config list; npm config get registry", "清华源 + npmmirror"),
    ("Q-03", "虚拟化已开启", "(Get-CimInstance Win32_Processor).VirtualizationFirmwareEnabled", "True"),
    ("Q-04", "WSL2 + Ubuntu 可用", "wsl -l -v", "Ubuntu-24.04 Running / VERSION 2"),
    ("Q-05", "Docker 可用", "docker run --rm hello-world", "输出 Hello from Docker!"),
    ("Q-06", "依赖服务全起", "docker compose -f docker-compose.dev.yml ps", "PG/Redis/EMQX/MinIO 均 healthy"),
    ("Q-07", "数据库可连", "docker exec campus-pg psql -U campus -d campus -c '\\dx'", "扩展含 vector"),
    ("Q-08", "后端可启", "curl -s http://localhost:8000/health", "status=ok"),
    ("Q-09", "接口文档可访问", "浏览器打开 http://localhost:8000/docs", "Swagger 正常"),
    ("Q-10", "前端可启", "浏览器访问 http://localhost:5173", "页面正常"),
    ("Q-11", "前后端联调通", "curl -s http://localhost:5173/api/health", "返回后端 JSON"),
    ("Q-12", "数据库迁移成功", "uv run alembic current", "显示版本号"),
    ("Q-13", "Git 仓库就绪", "git remote -v", "能看到远端地址"),
    ("Q-14", "环境自检通过", "python 环境工具\\scripts\\check-env.py --role dev", "核心项全部通过"),
    ("Q-15", "能完成一次提交", "git add . && git commit && git push", "远端可见提交记录"),
]
sheet(wb, "02-单机快速验收", "单机开发环境快速验收（15 项）",
      "对应文档 09 V2 第 4.6 节 ｜ 全部通过即可开始写功能代码",
      ["编号", "验收项", "校验命令", "通过标准", "结果", "备注", "验收人", "日期"],
      [list(q) + ["", "", "", ""] for q in QUICK],
      [8, 22, 46, 34, 8, 28, 10, 12], center_cols=(1, 5, 7, 8))

# ============================================================
# 03 全栈验收（25 项）
# ============================================================
FULL = [
    ("D-01", "WSL2 正常", "wsl -l -v", "Ubuntu-24.04 Running，VERSION=2"),
    ("D-02", "WSL 资源限制生效", "wsl -d Ubuntu-24.04 -- free -h", "内存接近设定值"),
    ("D-03", "apt 源为国内", "grep -r aliyun /etc/apt/sources.list.d/", "有 aliyun 记录"),
    ("D-04", "Docker 可用", "docker run --rm hello-world", "Hello from Docker!"),
    ("D-05", "Docker 加速生效", "docker info | Select-String 'Registry Mirrors'", "显示专属加速地址"),
    ("D-06", "PG16 + pgvector", "docker exec campus-pg psql -U campus -d campus -c '\\dx'", "含 vector 扩展"),
    ("D-07", "Redis 7 连通", "docker exec campus-redis redis-cli -a <口令> ping", "PONG"),
    ("D-08", "EMQX 连通", "curl -s http://localhost:18083/status", "返回状态信息"),
    ("D-09", "MinIO 存活", "curl -s http://localhost:9000/minio/health/live", "HTTP 200"),
    ("D-10", "后端依赖齐全", "uv run python -c \"import fastapi, sqlalchemy, asyncpg\"", "无报错"),
    ("D-11", "后端服务可启", "curl -s http://localhost:8000/health", "status=ok"),
    ("D-12", "Swagger 可访问", "浏览器打开 /docs", "接口文档正常"),
    ("D-13", "数据库迁移成功", "uv run alembic current", "显示最新版本号"),
    ("D-14", "NVIDIA 驱动正常", "nvidia-smi", "显示 RTX 3050"),
    ("D-15", "PyTorch CUDA 可用", "python -c \"import torch;print(torch.cuda.is_available())\"", "True"),
    ("D-16", "Ultralytics 可用", "yolo version", "输出版本号"),
    ("D-17", "最小训练跑通", "ls runs/detect/*/weights/best.pt", "权重文件存在"),
    ("D-18", "ONNX 导出成功", "ls runs/detect/*/weights/best.onnx", "文件存在"),
    ("D-19", "推理一致性达标", "对比脚本输出", "偏差 ≤1e-3"),
    ("D-20", "前端可启动", "浏览器访问 http://localhost:5173", "页面正常"),
    ("D-21", "前后端联调通", "curl -s http://localhost:5173/api/health", "返回后端 JSON"),
    ("D-22", "伪视频可拉流", "ffprobe rtsp://localhost:8554/cam1", "输出视频流信息"),
    ("D-23", "MQTT 模拟可上报", "MQTTX 订阅 iot/#", "持续收到消息"),
    ("D-24", "桩服务可访问", "curl -X POST http://localhost:9100/api/device/access/report", "返回模拟数据"),
    ("D-25", "环境自检通过", "python 环境工具\\scripts\\check-env.py --role dev", "通过率 100%"),
]
sheet(wb, "03-全栈验收", "台式机全栈环境验收（25 项）",
      "对应文档 09 V2 第 9 章 ｜ 全绿后可进入 M2 功能开发与联调",
      ["编号", "验收项", "校验命令", "通过标准", "结果", "实际输出/备注", "验收人", "日期"],
      [list(f) + ["", "", "", ""] for f in FULL],
      [8, 22, 46, 34, 8, 30, 10, 12], center_cols=(1, 5, 7, 8))

# ============================================================
# 04 可选环境验收（演示 / 协作 / 服务器 / 机房 / 现场）
# ============================================================
OPT = [
    ("演示机", "S-01", "依赖服务全起", "docker compose ps 全部 healthy"),
    ("演示机", "S-02", "后端 /health 正常", "status=ok"),
    ("演示机", "S-03", "前端可从 8080 访问", "页面正常且 /api 可用"),
    ("演示机", "S-04", "断网演示可用", "拔网线后完整演示无异常"),
    ("演示机", "S-05", "镜像已本地备份", "images-demo.tar 可导入"),
    ("演示机", "S-06", "演示录屏已备份", "录屏文件可播放"),
    ("协作同学", "C-01", "基础工具就绪", "git / node / pnpm 均有版本"),
    ("协作同学", "C-02", "镜像源为国内", "npm registry 为 npmmirror"),
    ("协作同学", "C-03", "代码可拉取", "git clone 成功且能 push"),
    ("协作同学", "C-04", "前端可启动", "pnpm dev 后页面正常"),
    ("协作同学", "C-05", "能连开发机后端", "页面可加载后端数据，无跨域报错"),
    ("协作同学", "C-06", "能完成一次提交", "远端可见其提交记录"),
    ("协作同学", "C-07", "已知晓协作红线", "确认不提交口令与人脸数据"),
    ("服务器", "K-01", "SSH 安全加固", "禁用 root 与口令登录"),
    ("服务器", "K-02", "时区与时间同步", "Asia/Shanghai 且 NTP 正常"),
    ("服务器", "K-03", "安全组最小开放", "仅 22/80/443（8883 限来源）"),
    ("服务器", "K-04", "Docker 与 Compose 可用", "版本命令正常"),
    ("服务器", "K-05", "依赖服务全起", "全部 healthy"),
    ("服务器", "K-06", "后端接口可访问", "/health 返回 ok"),
    ("服务器", "K-07", "HTTPS 有效", "无证书警告且可自动续期"),
    ("服务器", "K-08", "MQTT over TLS 可用", "校内网关可 TLS 上报"),
    ("服务器", "K-09", "自动备份生效", "备份文件生成且保留策略正确"),
    ("服务器", "K-10", "恢复演练通过", "恢复数据与源库一致"),
    ("服务器", "K-11", "监控告警可达", "手机收到阈值告警"),
    ("服务器", "K-12", "日志与磁盘", "日志按天切割，磁盘告警可用"),
    ("机房", "L-01", "母盘环境完成", "可独立完成一次开发与提交"),
    ("机房", "L-02", "批量部署完成", "抽查 5 台环境一致"),
    ("机房", "L-03", "离线包可用", "断网可安装依赖"),
    ("机房", "L-04", "集中开发服务器可用", "账号隔离，DB/API 可连"),
    ("机房", "L-05", "修复脚本可用", "5 分钟内恢复环境"),
    ("机房", "L-06", "学生账号隔离", "A 无法访问 B 的数据"),
    ("现场", "F-01", "平台服务全起", "容器 healthy，重启自恢复"),
    ("现场", "F-02", "基础数据导入完整", "组织/场所/人员/台账无重复"),
    ("现场", "F-03", "门禁联调", "刷卡刷脸 P95 ≤300ms"),
    ("现场", "F-04", "断网离线放行", "断网 30 分钟可放行且补传完整"),
    ("现场", "F-05", "视频接入", "预览延时 ≤1.5 秒"),
    ("现场", "F-06", "AI 算法达标", "准确率与误报率符合要求"),
    ("现场", "F-07", "通知通道", "短信/微信送达可查回执"),
    ("现场", "F-08", "定时任务", "归寝判定/报表/备份按计划执行"),
    ("现场", "F-09", "备份与回滚演练", "RPO ≤15min、RTO ≤2h、回滚 ≤15min"),
    ("现场", "F-10", "培训与文档移交", "三类培训签到，手册齐全"),
]
sheet(wb, "04-可选环境验收", "可选环境验收清单（演示 / 协作 / 服务器 / 机房 / 现场）",
      "这些环境都不是开发前置条件，按需配置后逐项打勾",
      ["环境", "编号", "验收项", "通过标准", "结果", "备注", "验收人", "日期"],
      [list(o) + ["", "", "", ""] for o in OPT],
      [10, 8, 26, 42, 8, 28, 10, 12], center_cols=(1, 2, 5, 7, 8))

wb.save(OUT)
print("已生成:", OUT)
print("工作表:", wb.sheetnames)
