# -*- coding: utf-8 -*-
import re
from docx import Document

d = Document(r"F:\职业中专\大赛\智慧园区创意\05_商业计划书.docx")
for x in d.paragraphs:
    if "**" in x.text:
        print("PARA:", repr(x.text[:120]))
        print("  runs:", [(r.text[:40], r.font.bold) for r in x.runs][:8])
        break
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            if "**" in c.text:
                print("CELL:", repr(c.text[:120]))
                print("  runs:", [(r.text[:40], r.font.bold) for p in c.paragraphs for r in p.runs][:8])
                raise SystemExit
