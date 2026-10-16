# -*- coding: utf-8 -*-
"""排查：09_新版 与 09_旧版 的文本提取差异，以及 RTX 3050 等的实际写法。"""
import os
from docx import Document

BASE = r"F:\职业中专\大赛\智慧园区创意\项目文档"

for name in ["09_电脑环境配置详细方案_新版.docx", "09_电脑环境配置详细方案.docx"]:
    p = os.path.join(BASE, name)
    d = Document(p)
    paras = [x.text for x in d.paragraphs]
    cells = []
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                cells.append(c.text)
    full = "\n".join(paras + cells)
    print(f"\n===== {name} =====")
    print("paragraphs:", len(paras), " table cells:", len(cells), " total chars:", len(full))
    for key in ["RTX 3050", "i5-13400F", "install-base.ps1", "wsl --install", "registry-mirrors",
                "pgvector", "torch.cuda.is_available", "Nginx", "内存 8GB"]:
        print(f"   {key:<26} {'√' if key in full else '×'}   出现次数 {full.count(key)}")
    print("   段落前 5 条:", [x[:40] for x in paras[:5]])
